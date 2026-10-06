"""Local document and chunk store used by ingestion and retrieval.

The bundled documents are small, attributed excerpts/notes from official public
sources, included only as a reproducible development corpus. They are not a
replacement for refreshing the official feeds.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import threading
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from app.models import Chunk, Document
from app.services.extraction import ExtractionUnavailable, extract as extract_document

INJECTION_PATTERNS = (
    r"ignore (?:all|any|previous) instructions",
    r"system prompt",
    r"developer message",
    r"reveal (?:your|the) (?:secret|prompt|system)",
    r"you are now",
    r"disregard (?:the|all)",
)


def _sha256(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _chunks_for(document: Document, chunk_size: int = 900) -> list[Chunk]:
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n|(?=^#{1,3} )", document.content, flags=re.MULTILINE) if part.strip()]
    if not paragraphs:
        paragraphs = [document.content.strip()]
    chunks: list[Chunk] = []
    ordinal = 0
    current_page = None
    current_section = ""
    for paragraph in paragraphs:
        page_match = re.match(r"^\[Page\s+(\d+)\]\s*", paragraph, flags=re.IGNORECASE)
        if page_match:
            current_page = int(page_match.group(1))
            paragraph = paragraph[page_match.end():].strip()
        heading = re.match(r"^#{1,6}\s+([^\n]+)", paragraph)
        if heading:
            current_section = heading.group(1).strip()
        words = paragraph.split()
        for start in range(0, len(words), chunk_size):
            text = " ".join(words[start : start + chunk_size]).strip()
            if not text:
                continue
            page_number = current_page
            page_match = re.match(r"^\[Page\s+(\d+)\]\s*", text, flags=re.IGNORECASE)
            if page_match:
                page_number = int(page_match.group(1))
                text = text[page_match.end():].strip()
            if not text:
                continue
            section = current_section
            if "[Table" in text or "[Table]" in text:
                section = "table"
            header = re.match(r"^#{1,3}\s+(.+?)(?:\s|$)", text)
            if header:
                section = header.group(1).strip()
            chunks.append(
                Chunk(
                    chunk_id=f"{document.document_id}:chunk:{ordinal}",
                    document_id=document.document_id,
                    text=text,
                    page_number=page_number,
                    section=section,
                    source=document.source,
                    publication_date=document.publication_date,
                    url=document.url,
                    version=document.version,
                    ordinal=ordinal,
                )
            )
            ordinal += 1
    return chunks


def scan_content(text: str) -> list[str]:
    """Return matched injection patterns; callers must decide whether to quarantine."""
    flags = []
    lowered = text.lower()
    for pattern in INJECTION_PATTERNS:
        if re.search(pattern, lowered):
            flags.append(pattern)
    return flags


# Kept private for compatibility with earlier imports.
_scan_content = scan_content


class DocumentStore:
    def __init__(self, *, include_bootstrap: bool = False) -> None:
        self.lock = threading.RLock()
        self.include_bootstrap = include_bootstrap
        self._bootstrap_stamp = None
        self.documents: dict[str, Document] = {}
        self.chunks: dict[str, Chunk] = {}
        self.chunk_counts: dict[str, int] = {}
        self.version = 0
        self._load_seed()
        if include_bootstrap:
            self._load_bootstrap()
            self._bootstrap_stamp = self._bootstrap_signature()

    @staticmethod
    def _bootstrap_signature():
        from app.config import settings
        path = settings.data_dir / "normalized/cyber_entities.jsonl"
        return (path.stat().st_mtime_ns, path.stat().st_size) if path.exists() else None

    def refresh_bootstrap(self) -> None:
        if not self.include_bootstrap:
            return
        signature = self._bootstrap_signature()
        if signature is None or signature == self._bootstrap_stamp:
            return
        with self.lock:
            if signature == self._bootstrap_stamp:
                return
            refreshed = DocumentStore(include_bootstrap=True)
            # Retain owned uploads while replacing the shared source snapshot.
            retained = [doc for doc in self.documents.values() if "bootstrap" not in doc.tags]
            for doc in retained:
                refreshed.add_document(doc)
            self.documents, self.chunks, self.chunk_counts = refreshed.documents, refreshed.chunks, refreshed.chunk_counts
            self.version += 1
            self._bootstrap_stamp = signature

    def _load_seed(self) -> None:
        for document in seed_documents():
            self.add_document(document)

    def _load_bootstrap(self) -> None:
        """Load only successfully normalized bootstrap records as shared knowledge."""
        try:
            from app.config import settings
            path = settings.data_dir / "normalized" / "cyber_entities.jsonl"
            if not path.exists():
                return
            for line in path.read_text(encoding="utf-8").splitlines():
                if not line.strip():
                    continue
                row = json.loads(line)
                source = str(row.get("source", "bootstrap"))
                entity_type = str(row.get("entity_type", "UNKNOWN"))
                refs = row.get("references") if isinstance(row.get("references"), list) else []
                metadata = row.get("metadata") if isinstance(row.get("metadata"), dict) else {}
                url = next((str(ref.get("url")) for ref in refs if isinstance(ref, dict) and str(ref.get("url", "")).startswith(("http://", "https://"))), "")
                if not url:
                    url = next((str(value) for value in metadata.values() if isinstance(value, str) and value.startswith(("http://", "https://"))), "")
                raw_id = str(row.get("id") or row.get("source_id") or "")
                if not raw_id:
                    continue
                metadata_lines = []
                if entity_type == "CVE" or source == "CISA KEV":
                    for key in ("cvss", "products", "weaknesses", "vendorProject", "product", "dateAdded", "dueDate", "requiredAction", "knownRansomwareCampaignUse", "notes", "vulnStatus"):
                        if key in metadata and metadata[key] not in (None, "", [], {}):
                            metadata_lines.append(f"{key}: {json.dumps(metadata[key], ensure_ascii=False)}")
                content = "\n".join(part for part in (f"{row.get('name', '')} [{raw_id}]", str(row.get('description', '')), *metadata_lines) if part).strip()
                authority = {"CISA": 0.98, "NVD": 0.96, "MITRE ATT&CK": 0.97, "MITRE ATLAS": 0.96, "CWE": 0.94, "CAPEC": 0.94, "NIST": 0.96, "OWASP": 0.92}.get(source, 0.75)
                document = Document(
                    document_id=f"kb-{hashlib.sha256(f'{source}:{raw_id}'.encode()).hexdigest()[:24]}",
                    title=str(row.get("name") or raw_id),
                    source=source,
                    source_type="paper" if entity_type == "RESEARCH_PAPER" else "authoritative_dataset",
                    url=url,
                    publication_date=str(row.get("created_at") or ""),
                    update_date=str(row.get("updated_at") or ""),
                    version=str(row.get("version") or ""),
                    authority_score=authority,
                    tags=["bootstrap", source.lower().replace(" ", "_"), entity_type.lower()],
                    checksum=hashlib.sha256(content.encode()).hexdigest(),
                    content=content,
                )
                self.add_document(document)
        except (OSError, json.JSONDecodeError):
            # Missing or malformed optional bootstrap output leaves the curated
            # corpus available; the manifest records the acquisition failure.
            return

    def add_document(self, document: Document) -> Document:
        with self.lock:
            return self._add_document(document)

    def _add_document(self, document: Document) -> Document:
        document.security_flags = scan_content(document.content)
        was_present = document.document_id in self.documents
        self.documents[document.document_id] = document
        if was_present:
            for chunk_id in [key for key, value in self.chunks.items() if value.document_id == document.document_id]:
                del self.chunks[chunk_id]
        self.chunk_counts[document.document_id] = 0
        # A flagged document is retained for audit/deletion but never enters
        # retrieval until an explicit, reviewed re-index operation exists.
        if not document.security_flags:
            for chunk in _chunks_for(document):
                self.chunks[chunk.chunk_id] = chunk
                self.chunk_counts[document.document_id] += 1
        self.version += 1
        return document

    def load_persisted(self, record: dict[str, object], owner_id: str) -> Document | None:
        """Restore one owned document from repository metadata without widening visibility."""
        metadata = record.get("metadata")
        if not isinstance(metadata, dict):
            return None
        content = metadata.get("content")
        if not isinstance(content, str):
            return None
        document = Document(
            document_id=str(record.get("id") or metadata.get("document_id") or ""),
            title=str(metadata.get("title") or record.get("title") or "Untitled"),
            source=str(metadata.get("source") or record.get("source") or "user_upload"),
            source_type=str(metadata.get("source_type") or "user_document"),
            url=str(metadata.get("url") or ""),
            author=str(metadata.get("author") or ""),
            publication_date=str(metadata.get("publication_date") or ""),
            update_date=str(metadata.get("update_date") or ""),
            version=str(metadata.get("version") or ""),
            authority_score=float(metadata.get("authority_score") or 0.35),
            language=str(metadata.get("language") or "en"),
            tags=[str(tag) for tag in (metadata.get("tags") or []) if str(tag) != f"owner:{owner_id}"],
            checksum=str(metadata.get("checksum") or record.get("checksum") or ""),
            ingestion_timestamp=str(metadata.get("ingestion_timestamp") or record.get("created_at") or datetime.now(timezone.utc).isoformat()),
            content=content,
            security_flags=[str(flag) for flag in (metadata.get("security_flags") or [])],
            extraction=dict(metadata.get("extraction") or {}),
        )
        if f"owner:{owner_id}" not in document.tags:
            document.tags.append(f"owner:{owner_id}")
        self.add_document(document)
        return document

    def ingest_bytes(
        self,
        filename: str,
        payload: bytes,
        *,
        source: str = "user_upload",
        source_type: str = "user_document",
        url: str = "",
        owner_id: str = "local",
    ) -> Document:
        checksum = _sha256(payload)
        owner_key = hashlib.sha256(owner_id.encode("utf-8")).hexdigest()[:12]
        document_id = f"doc-{owner_key}-{checksum[:16]}"
        suffix = Path(filename).suffix.lower()
        extraction = extract_document(filename, payload)
        content = extraction.text
        title = Path(filename).stem.replace("_", " ").replace("-", " ").strip() or filename
        document = Document(
            document_id=document_id,
            title=title,
            source=source,
            source_type=source_type,
            url=url,
            authority_score=0.35,
            checksum=checksum,
            content=content,
            tags=[suffix.lstrip(".") or "text", f"owner:{owner_id}"],
            extraction={"engine": extraction.engine, "spans": extraction.spans, "warnings": extraction.warnings},
        )
        return self.add_document(document)

    @staticmethod
    def _extract_content(suffix: str, payload: bytes) -> str:
        if suffix == ".json":
            try:
                value = json.loads(payload.decode("utf-8", errors="replace"))
                return json.dumps(value, indent=2, ensure_ascii=False)
            except json.JSONDecodeError:
                return payload.decode("utf-8", errors="replace")
        if suffix == ".csv":
            rows = csv.reader(payload.decode("utf-8", errors="replace").splitlines())
            return "\n".join(" | ".join(row) for row in rows)
        if suffix == ".pdf":
            try:
                import fitz  # type: ignore
                pages: list[str] = []
                with fitz.open(stream=payload, filetype="pdf") as pdf:
                    for page_number, page in enumerate(pdf, start=1):
                        text = page.get_text("text").strip()
                        if not text:
                            try:
                                from PIL import Image  # type: ignore
                                import pytesseract  # type: ignore
                                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                                image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
                                text = pytesseract.image_to_string(image).strip()
                            except ImportError:
                                text = "OCR unavailable: install optional Pillow, pytesseract, and the tesseract binary."
                        table_text = ""
                        try:
                            tables = page.find_tables()
                            rows = []
                            for table in tables.tables:
                                rows.extend(" | ".join(str(cell or "") for cell in row) for row in table.extract())
                            table_text = "\n[Table]\n" + "\n".join(rows) if rows else ""
                        except Exception:
                            table_text = ""
                        pages.append(f"[Page {page_number}]\n{text}{table_text}".strip())
                return "\n\n".join(pages)
            except Exception as exc:
                return f"PDF extraction unavailable: {exc}. The file is stored but requires the optional pymupdf dependency."
        if suffix == ".docx":
            try:
                from docx import Document as DocxDocument  # type: ignore
                document = DocxDocument(__import__("io").BytesIO(payload))
                parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
                for index, table in enumerate(document.tables, start=1):
                    parts.append(f"[Table {index}]\n" + "\n".join(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows))
                return "\n\n".join(parts)
            except ImportError:
                return "DOCX extraction unavailable: install optional python-docx."
            except Exception as exc:
                return f"DOCX extraction failed: {exc}"
        if suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif"}:
            try:
                from PIL import Image  # type: ignore
                import pytesseract  # type: ignore
                image = Image.open(__import__("io").BytesIO(payload))
                text = pytesseract.image_to_string(image).strip()
                return f"[Image OCR]\n{text}" if text else "Image contained no OCR text. Vision interpretation is unavailable; no visual cybersecurity claim was generated."
            except ImportError:
                return "Image uploaded. Vision interpretation is unavailable in the deterministic local mode; install optional Pillow, pytesseract, and tesseract for OCR. No visual cybersecurity claim was generated."
            except Exception as exc:
                return f"Image OCR failed: {exc}. No visual cybersecurity claim was generated."
        return payload.decode("utf-8", errors="replace")

    def list_documents(self, owner_id: str | None = None, *, include_shared: bool = True) -> list[Document]:
        self.refresh_bootstrap()
        with self.lock:
            documents = list(self.documents.values())
        if owner_id:
            owner_tag = f"owner:{owner_id}"
            documents = [doc for doc in documents if owner_tag in doc.tags or (include_shared and not any(tag.startswith("owner:") for tag in doc.tags))]
        return sorted(documents, key=lambda doc: doc.ingestion_timestamp, reverse=True)

    def visible_document_ids(self, owner_id: str) -> set[str]:
        return {doc.document_id for doc in self.list_documents(owner_id, include_shared=True)}

    def delete_document(self, document_id: str, owner_id: str) -> bool:
        document = self.documents.get(document_id)
        if document is None or f"owner:{owner_id}" not in document.tags:
            return False
        self.documents.pop(document_id, None)
        self.chunk_counts.pop(document_id, None)
        for chunk_id in [key for key, value in self.chunks.items() if value.document_id == document_id]:
            del self.chunks[chunk_id]
        self.version += 1
        return True

    def get_document(self, document_id: str, owner_id: str | None = None) -> Document | None:
        document = self.documents.get(document_id)
        if document is None or (owner_id and f"owner:{owner_id}" not in document.tags and any(tag.startswith("owner:") for tag in document.tags)):
            return None
        return document

    def chunk_count(self, document_id: str) -> int:
        return self.chunk_counts.get(document_id, 0)

    def get_chunks(self, document_ids: set[str] | None = None) -> list[Chunk]:
        values = list(self.chunks.values())
        if document_ids is not None:
            values = [chunk for chunk in values if chunk.document_id in document_ids]
        return values



# Public sources used for the reproducible seed corpus. Statements are concise
# and intentionally link back to the official page for inspection.
def seed_documents() -> Iterable[Document]:
    timestamp = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    yield Document(
        document_id="seed-nist-csf-2-0",
        title="NIST Cybersecurity Framework 2.0",
        source="NIST",
        source_type="official_framework",
        url="https://www.nist.gov/cyberframework",
        author="National Institute of Standards and Technology",
        publication_date="2024-02-26",
        update_date="2024-02-26",
        version="2.0",
        authority_score=0.98,
        tags=["framework", "risk-management", "authoritative"],
        ingestion_timestamp=timestamp,
        content=(
            "# Overview\nThe NIST Cybersecurity Framework (CSF) 2.0 provides guidance to help organizations manage cybersecurity risks. "
            "It is organized around the Functions Govern, Identify, Protect, Detect, Respond, and Recover.\n\n"
            "# Evidence note\nThe CSF is a risk-management framework. It does not prescribe one technology or guarantee that an organization is secure. "
            "Organizations should use organizational context and risk outcomes when applying the framework."
        ),
    )
    yield Document(
        document_id="seed-mitre-attack-t1059",
        title="MITRE ATT&CK T1059 Command and Scripting Interpreter",
        source="MITRE ATT&CK",
        source_type="official_knowledge_base",
        url="https://attack.mitre.org/techniques/T1059/",
        author="MITRE",
        publication_date="",
        update_date="",
        version="enterprise",
        authority_score=0.97,
        tags=["attack", "technique", "detection", "authoritative"],
        ingestion_timestamp=timestamp,
        content=(
            "# Technique\nMITRE ATT&CK technique T1059 is Command and Scripting Interpreter. The technique describes adversaries using command "
            "and scripting interpreters to execute commands, scripts, or binaries. ATT&CK lists platform-specific sub-techniques and "
            "provides procedure examples, detection guidance, and mitigations on the technique page.\n\n"
            "# Interpretation\nA log containing a command interpreter event is a lead for investigation, not proof by itself that an intrusion occurred. "
            "Detection should be correlated with user, process, parent process, host, timing, and surrounding activity."
        ),
    )
    yield Document(
        document_id="seed-cisa-kev",
        title="CISA Known Exploited Vulnerabilities Catalog",
        source="CISA",
        source_type="official_catalog",
        url="https://www.cisa.gov/known-exploited-vulnerabilities-catalog",
        author="Cybersecurity and Infrastructure Security Agency",
        publication_date="",
        update_date="",
        version="current catalog",
        authority_score=0.98,
        tags=["cve", "kev", "vulnerability", "authoritative"],
        ingestion_timestamp=timestamp,
        content=(
            "# Catalog meaning\nThe CISA Known Exploited Vulnerabilities (KEV) Catalog is an authoritative source of vulnerabilities that CISA says "
            "have been exploited in the wild. The catalog includes fields such as the vulnerability identifier, vendor, product, due date, "
            "and required action for federal civilian agencies.\n\n"
            "# Interpretation\nAbsence from the KEV Catalog should not be interpreted as evidence that a vulnerability cannot be exploited. "
            "Catalog status and due dates should be checked against the current official catalog."
        ),
    )
    yield Document(
        document_id="seed-nvd-cve-records",
        title="NVD CVE and CVSS data guidance",
        source="NVD",
        source_type="official_vulnerability_database",
        url="https://nvd.nist.gov/developers/vulnerabilities",
        author="National Vulnerability Database",
        publication_date="",
        update_date="",
        version="API v2.0 documentation",
        authority_score=0.96,
        tags=["cve", "cvss", "vulnerability", "authoritative"],
        ingestion_timestamp=timestamp,
        content=(
            "# API and records\nThe National Vulnerability Database publishes machine-readable vulnerability information and documents a version 2.0 API. "
            "A CVE record can contain descriptions, configurations, references, weakness information, and CVSS data when available.\n\n"
            "# Interpretation\nA CVSS score is a severity measure with a defined scoring specification; it is not a prediction of exploit probability or business risk. "
            "Consumers should inspect the record's source, version, timestamps, and references rather than copying a score without provenance."
        ),
    )
    yield Document(
        document_id="seed-owasp-llm-prompt-injection",
        title="OWASP Top 10 for Large Language Model Applications",
        source="OWASP",
        source_type="official_project",
        url="https://owasp.org/www-project-top-10-for-large-language-model-applications/",
        author="OWASP Foundation",
        publication_date="",
        update_date="",
        version="2025",
        authority_score=0.94,
        tags=["llm", "prompt-injection", "rag-security", "authoritative"],
        ingestion_timestamp=timestamp,
        content=(
            "# Prompt injection\nPrompt injection occurs when an attacker crafts input that causes a language model to perform unintended actions or follow "
            "untrusted instructions. In retrieval-augmented systems, retrieved content must be treated as data and must not override the system "
            "or developer policy.\n\n# Defensive interpretation\nUse content isolation, least privilege, validation, monitoring, and explicit trust boundaries. "
            "A security review should distinguish a document's factual claims from instructions embedded in that document."
        ),
    )
