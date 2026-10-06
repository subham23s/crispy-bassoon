"""Generate an evidence-linked Word work record and optional verification logs.

The narrative distinguishes earlier repository records from current execution.
The function appendix is generated from source, not reconstructed chat history.
"""
from __future__ import annotations

import argparse
import ast
import hashlib
import json
import os
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/reports"


def read_json(relative: str, default=None):
    path = ROOT / relative
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {} if default is None else default


def inventory():
    files = []
    for directory in ("backend", "scripts", "ml", "frontend/app", "frontend/components", "frontend/lib", "frontend/services", "frontend/hooks", "frontend/types", "frontend/tests", "mobile/lib", "mobile/test"):
        for path in sorted((ROOT / directory).rglob("*")):
            if path.suffix in {".py", ".ts", ".tsx", ".dart", ".mjs"} and not any(part in {"__pycache__", "node_modules", ".next"} for part in path.parts):
                files.append(path)
    return sorted(set(files))


def python_definitions(path: Path):
    tree = ast.parse(path.read_text(encoding="utf-8"))
    found = []
    def visit(node, prefix=""):
        for child in ast.iter_child_nodes(node):
            if isinstance(child, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                name = prefix + child.name
                docstring = ast.get_docstring(child) or ""
                kind = "class" if isinstance(child, ast.ClassDef) else "async function" if isinstance(child, ast.AsyncFunctionDef) else "function/method"
                signature = name if isinstance(child, ast.ClassDef) else name + "(" + ast.unparse(child.args) + ")"
                found.append((child.lineno, kind, signature, docstring.split("\n")[0]))
                visit(child, name + ".")
            else:
                visit(child, prefix)
    visit(tree)
    return ast.get_docstring(tree) or "", found


def run_checks():
    commands = [
        ("Backend regression suite", [sys.executable, "-m", "pytest", "-p", "no:cacheprovider", "-q", "backend/tests"], 240),
        ("Python compilation", [sys.executable, "-m", "compileall", "-q", "backend", "scripts", "ml"], 120),
        ("Web client tests", ["npm", "test", "--prefix", "frontend"], 120),
        ("TypeScript typecheck", ["npm", "run", "typecheck", "--prefix", "frontend"], 120),
        ("Next.js production build", ["npm", "run", "build", "--prefix", "frontend"], 240),
        ("Docker configuration checker", [sys.executable, "scripts/check_docker_config.py"], 60),
        ("Mobile environment checker", [sys.executable, "scripts/check_mobile.py", "--json"], 60),
        ("Environment availability", [sys.executable, "scripts/check_environment.py"], 90),
    ]
    results = []
    for label, command, timeout in commands:
        started = datetime.now(timezone.utc).isoformat()
        try:
            result = subprocess.run(command, cwd=ROOT, env={**os.environ, "PYTHONPATH": str(ROOT / "backend") + os.pathsep + str(ROOT), "NEXT_TELEMETRY_DISABLED": "1"}, capture_output=True, text=True, timeout=timeout)
            row = {"check": label, "command": subprocess.list2cmdline(command), "started_at": started, "finished_at": datetime.now(timezone.utc).isoformat(), "exit_code": result.returncode, "status": "PASSED" if result.returncode == 0 else "FAILED / BLOCKED", "output": result.stdout + result.stderr}
        except (OSError, subprocess.TimeoutExpired) as exc:
            row = {"check": label, "command": subprocess.list2cmdline(command), "started_at": started, "finished_at": datetime.now(timezone.utc).isoformat(), "exit_code": None, "status": "BLOCKED / TIMEOUT", "output": str(exc)}
        results.append(row)
        print(f"{label}: {row['status']}", flush=True)
    return results


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--verify", action="store_true", help="Execute checks and embed their actual output")
    args = parser.parse_args()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    verification_path = OUTPUT / "work_record_verification.json"
    if args.verify:
        checks = run_checks()
        verification_path.write_text(json.dumps(checks, indent=2), encoding="utf-8")
    else:
        checks = read_json("docs/reports/work_record_verification.json", [])
    generated_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("17365D")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("CyberSecRA-X work record | Page ")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    doc.core_properties.title = "CyberSecRA-X: Complete Work Record, Implementation Inventory and Remaining Work"
    doc.core_properties.subject = "Evidence-linked project handover covering 2–3 October 2026"
    doc.core_properties.author = "CyberSecRA-X project work record"

    def heading(text, level=1): doc.add_heading(text, level)
    def para(text): doc.add_paragraph(text)
    def bullets(items):
        for item in items: doc.add_paragraph(item, style="List Bullet")
    def table(headers, rows):
        value = doc.add_table(rows=1, cols=len(headers)); value.style = "Light Shading Accent 1"
        for cell, text in zip(value.rows[0].cells, headers): cell.text = str(text)
        header = OxmlElement("w:tblHeader"); value.rows[0]._tr.get_or_add_trPr().append(header)
        for row in rows:
            cells = value.add_row().cells
            for cell, text in zip(cells, row): cell.text = str(text)
        doc.add_paragraph()
    def code(text):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(text); run.font.name = "Consolas"; run.font.size = Pt(8)

    doc.add_heading("CyberSecRA-X", 0)
    doc.add_heading("Complete Work Record & Technical Handover", 1)
    para("Evidence-Aware Cybersecurity Research Intelligence")
    para("Coverage: project inception recorded on 2 October 2026 through the source snapshot on 3 October 2026.")
    para(f"Document generated (UTC): {generated_at}")
    para("Prepared as a detailed record for explaining the project to a supervisor, examiner, teammate, or future maintainer.")
    table(["Record scope", "Meaning"], [
        ("Earlier recorded work", "Derived from README, CHANGELOG, saved agent plan, quality gate, architecture decisions, final report, and stored data/evaluation artifacts."),
        ("Current-session work", "Changes implemented during this resumed session and commands executed in it. Some recent paths still require targeted regression tests and live integration checks."),
        ("Minute-by-minute history", "A complete timestamped transcript is unavailable and the Git repository has no commits. Exact times are only given where an artifact or a captured check records them. The function appendix supplies granular source references instead of invented activity timestamps."),
        ("Completion boundary", "This is a substantial local-first implementation with real source artifacts and optional adapters. It is not a declaration that all production infrastructure, transformer models, mobile builds, or scientific validation are complete."),
    ])
    heading("How to use this document")
    bullets(["Sections 1–5 explain the purpose, history, architecture, and user-facing workflows.", "Sections 6–12 describe implementation details, data, ML, CERS, security, and environment setup.", "Sections 13–16 record verification, remaining work, demonstration steps, and explanation notes.", "Appendices enumerate API routes, repository tables, Python classes/functions, frontend/mobile modules, exact check outputs, and source references."])
    doc.add_page_break()

    heading("1. Project purpose and explanation")
    para("CyberSecRA-X is a defensive cybersecurity research platform built around the question: ‘Why should I trust this answer?’ Its core output connects a question to attributed evidence, claim-level verification, possible contradictions, citations, provenance, a project-defined reliability score, and explicit abstention when the evidence is insufficient.")
    bullets(["Supported subject areas include vulnerability intelligence, threat intelligence, cybersecurity research, defensive log/IOC triage, research-paper comparison, and research-gap exploration.", "The product includes a Python backend, a Next.js/TypeScript browser application, a Flutter mobile client, data acquisition/processing scripts, evaluation tooling, and deployment configuration.", "Its defining workflow is: understand → retrieve → rank → correlate → generate or summarize → verify → detect conflicts → score → cite → abstain if necessary.", "A deterministic evidence summary is available without an LLM. A configured provider can stream actual generated text; generated text remains provisional until final verification completes.", "The corpus distinguishes source snapshots, uploaded documents, manually curated supervision, development benchmarks, and optional synthetic test/training examples. No new factual dataset or human calibration labels were created in this resumed session."])

    heading("2. Reconstructed implementation history")
    plan = read_json("__agent__/plan.json")
    table(["Stage / recorded time", "Work recorded", "Evidence / qualification"], [
        ("Initial increment — 2026-10-02", "Monorepo, local evidence pipeline, query understanding, OOD detection, hybrid retrieval, claims, CERS, auth, SQLite, uploads, RAG security, log/IOC analysis, web/mobile clients, evaluation, deployment docs.", "CHANGELOG.md; source files. Initial date is recorded, not a reconstructed minute-by-minute log."),
        ("Data/bootstrap expansion — prior session", "Official-source ingestion; source manifests; normalization; entities/relationships; chunks; BM25; feature-hash embeddings; persistent local vector artifact; baseline ML training.", "docs/FINAL_PROJECT_REPORT.md and data artifacts. Earlier reported counts appear in Section 7."),
        ("Application expansion — prior session", "Production adapters; NLI/reranker adapters; research workspace; Next.js routes; graph/timeline/admin/experiments; report export; background upload manager; mobile API flows.", "Saved plan and current source. Implemented source does not establish live execution of every adapter."),
        (f"Saved checkpoint — {plan.get('updated_at', 'unrecorded')}", "Nine saved plan items marked done, including the final reporting task.", "__agent__/plan.json. This marked the preceding increment complete, not external integrations or scientific validation."),
        ("Resume — 2026-10-03", "Recovered source/docs/plan; found no Git commits; ran baseline tests; re-used /tmp/cybersecra-venv because system Python lacks pip; identified real implementation gaps.", "Current session tool output; no source rollback or project restart."),
        ("Implementation during this session", "Graph BFS fixes; retrieval optimization and real graph links; shared sync/SSE pipeline; answer/chunk persistence; export snapshots; workspace/paper isolation; paper analysis/comparison; durable upload spool; calibration fitter; independent ablation switches; scheduler/reindex worker; client streaming/report workflows; environment repairs.", "Detailed current-session change register in Section 5; current source inventory in appendices."),
    ])
    heading("Saved prior-session plan", 2)
    table(["Task recorded in plan", "Saved status"], [(item["title"], item["status"]) for item in plan.get("items", [])])

    heading("3. Architecture and repository layout")
    para("Web/mobile clients call the same authenticated API. Runtime initializes infrastructure clients, the selected repository, authentication, the document store, the query pipeline, the graph service, refresh jobs, and document-processing jobs. SQLite is the local default; configured PostgreSQL, Redis, and Neo4j are probed and reported individually.")
    code("Browser / Flutter → HTTP API → Auth + owner checks → Query pipeline\n                              ↓                       ↓\n                       Repository / jobs       Evidence + verification\n                              ↑                       ↑\nOfficial feeds → normalize → chunks / indices / graph → retrieval\n                              ↓\n                Evaluation + calibration + report exports")
    table(["Location", "Purpose"], [
        ("backend/app/main.py", "Runtime, REST routing, HTTP/SSE transport, ownership checks, workspaces, papers, exports, fallback server."),
        ("backend/app/services/", "Corpus, retrieval, pipeline, LLM provider, extraction, jobs, graph, reranking, refresh and security services."),
        ("backend/app/database/", "Repository schema/methods; SQLite and PostgreSQL; optional Redis and Neo4j adapters; infrastructure health."),
        ("backend/app/verification/", "Claim extraction, NLI adapter/fallback, citation verification, contradictions, CERS and human-label calibration."),
        ("backend/app/research/", "Paper fields/comparison, candidate gap analysis, and evaluation metrics."),
        ("scripts/data/; scripts/evaluation/; ml/", "Acquisition/normalization/indexing; dataset validation/training; benchmark/ablation/calibration runners; inference adapters."),
        ("frontend/app/; frontend/components/; frontend/lib/", "Next.js routes, interactive dashboards/workspaces, typed API/auth/streaming client."),
        ("mobile/lib/", "Flutter authentication, query/voice, uploads/camera, defensive tools, research views, secure token handling."),
        ("data/; docs/; docker/; .github/", "Real source artifacts and evaluation outputs; documentation; container initialization; CI."),
    ])

    heading("4. Functional workflows and their major operations")
    workflows = [
        ("Authentication", ["Register or log in with email/password; hash passwords using salted PBKDF2.", "Issue signed expiring access tokens and rotate persisted refresh tokens.", "Protected routes authenticate the bearer token; admin routes additionally inspect role.", "Browser and mobile clients retry authentication once after refresh; mobile stores tokens in secure storage."]),
        ("Evidence question", ["Validate question and optional owned workspace; classify intent/domain/language/entities and temporal/multihop needs.", "Select visible shared/owned evidence, exclude quarantined content, and perform configured retrieval/ranking.", "Check requested identifiers against the retrieved evidence; abstain when identifiers/evidence are missing.", "Generate using the configured LLM stream or produce an evidence-based deterministic summary.", "Verify claims and citation markers; attach source conflicts; compute CERS and abstention.", "Persist the completed answer and optionally its workspace research item; return the same final contract through ordinary HTTP or SSE."]),
        ("Upload and ingestion", ["Accept supported document bytes, validate base64 and size, and extract real content.", "Support UTF-8 text/Markdown/logs, JSON, CSV, PDF text/tables, DOCX, and OCR when available.", "Compute a checksum and owner-specific document ID; keep extraction spans/warnings.", "Scan content for prompt-injection signals; quarantine flagged documents rather than retrieving them.", "Create chunks with provenance; persist the document and normalized chunks.", "For background work, persist the input and job metadata, process through bounded workers, expose lifecycle state, and recover accepted jobs after restart."]),
        ("Research workspace", ["Create and list owned workspaces; add questions/research items and notes.", "Validate workspace, source-document, and referenced research-item ownership.", "Ask and save a completed answer to the selected workspace.", "Save an actual visible evidence chunk as a server-side snapshot; list workspace evidence."]),
        ("Paper intelligence", ["Upload a document tagged as a paper or inspect permitted bootstrap paper metadata.", "Extract recognizable sections, reported metrics/results/limitations, references and source locations.", "Leave missing fields explicit; restrict year extraction to header metadata rather than arbitrary reference years.", "Analyze one visible, non-quarantined document or compare two or more papers across extracted fields.", "Generate candidate research-gap observations from supplied evidence; these remain candidates, not proven novelty claims."]),
        ("Graph and timeline", ["Search stored normalized graph nodes/edges with provenance or use a configured Neo4j graph.", "Apply bounded breadth-first traversal; deduplicate edges and return endpoints that exist in the returned subgraph.", "Use actual neighboring graph documents/relationships as optional retrieval signals.", "List publication/update/refresh timeline events from available metadata, with date filters."]),
        ("Reports", ["Select JSON, Markdown, DOCX or PDF.", "Export a stored answer by answer_id so evidence and text match the viewed response.", "Include claims, support, citation IDs, provenance, conflicts, graph links, timeline and limitations.", "Wrap/paginate PDF text rather than truncating each line; optional exporters report missing dependencies."]),
        ("Defensive tools and evaluation", ["Classify supplied log/IOC text conservatively; local heuristics are labeled and do not establish maliciousness.", "Run RAG-security pattern scans and persisted-document assessments.", "Create experiments and store their metrics; inspect computed benchmark artifacts.", "Run independently configured ablations and fit experimental calibration only from complete attributed human records with held-out evaluation data."]),
    ]
    for title, details in workflows: heading(title, 2); bullets(details)

    heading("5. Detailed current-session implementation register")
    changes = [
        ("Graph traversal and Neo4j", "backend/app/services/graph.py; database/neo4j_adapter.py; scripts/sync_graph.py", "Fixed same-loop cascading traversal; bounded depth/limit; exact matches first; closed subgraphs; deduplicated edges; explicit live BFS; idempotent batched import of normalized nodes/resolved edges.", "Local graph regression existed; new boundary and live-import checks still needed. Neo4j service unavailable."),
        ("Retrieval", "backend/app/services/retrieval.py", "Added corpus-IDF BM25, cached term counts/norms/postings, early candidate filtering, independent ranking switches and normalized component weights; links actual graph neighbors to corpus document IDs.", "Existing pipeline tests passed after the changes; large-corpus speed and ranking quality require dedicated measurement."),
        ("Shared query execution", "backend/app/services/pipeline.py", "Introduced PipelineOptions; ordinary ask consumes the same execution stream as SSE; unified verification/abstention; standardized citation labels; discarded interrupted provider output before deterministic replacement; require all requested identifiers.", "Deterministic path exercised by tests; external provider and model-backed paths remain unverified."),
        ("Answer and chunk persistence", "backend/app/database/{sqlite,postgres,interfaces}.py; main.py; models/domain.py", "Added owned answer snapshots/history and chunk persistence/access methods; SQLite document/workspace upserts avoid delete-and-reinsert side effects; restore serialized answers for exports; normalized PostgreSQL URLs and some datetime serialization.", "SQLite/API regressions passed at checkpoints; new PostgreSQL methods need a live service and new targeted tests."),
        ("Workspace and paper boundaries", "backend/app/main.py", "Validate owned workspace/source/item references; save actual visible chunk snapshots; restrict paper/threat/vulnerability listings to visible documents; exclude quarantined content from gaps/papers/timeline; add paper analyze/compare and answer-history routes.", "Source implemented; targeted cross-user regressions for new routes still to be completed."),
        ("Paper extraction", "backend/app/research/papers.py", "Preserve section line/page locations and reference line breaks; inspect header year metadata; extract metrics from evaluation/results sections; attach document IDs and URLs through paper APIs.", "Existing paper tests passed after changes; extraction quality on diverse real papers remains a task."),
        ("Multimodal chunk provenance", "backend/app/services/corpus.py", "Restore extraction metadata on restart; propagate page/section state across chunks; refresh changed shared bootstrap snapshots while preserving uploads.", "Text/ownership paths tested by existing suite; extensive PDF/OCR/DOCX/refresh concurrency tests remain."),
        ("Durable background uploads", "backend/app/services/document_jobs.py; main.py", "Bounded accepted queue; atomic input/job spool; restart recovery; persistent endpoint fallback; completed indexed/quarantined status; worker shutdown and logged repository update failures.", "Existing lifecycle test passed at checkpoint; restart/fault/multi-process tests still needed."),
        ("Citation verification", "backend/app/verification/citation_verifier.py", "Recognize C/E markers, verify against referenced evidence, reject out-of-range citations, and avoid heuristic partial-support upgrades when explicit transformer verification is unavailable.", "Existing verification/pipeline tests passed; new adversarial marker tests remain."),
        ("NLI model loading", "backend/app/verification/nli_verifier.py; scripts/setup_models.py", "Default inference loads cached models only; explicit setup permits downloads; pair input truncated to supported length.", "Optional model weights not downloaded; transformer inference not claimed."),
        ("CERS calibration", "backend/app/verification/calibration.py; cers.py; scripts/evaluation/calibrate_cers.py", "Input validation, annotator/split requirements, duplicate rejection, isotonic regression, empirical risk threshold selection and held-out Brier/ECE/coverage/risk; optional fitted artifact integration bound to weights and label checksum.", "Implementation exists; no human labels fitted and no calibration metrics claimed. Targeted numeric/edge-case tests pending."),
        ("Independent ablations", "scripts/evaluation/run_ablation.py", "B lexical; C hybrid baseline; D–I each add one component; J all configured components; same deterministic generator and shared corpus/index; proper recall denominator and reciprocal rank; retained answer snapshots and latencies.", "New runner not yet executed at report-generation start. Old artifact results do not validate the revised runner."),
        ("Source refresh and workers", "services/source_refresh.py; jobs/refresh.py; scripts/refresh_worker.py; scripts/data/common.py", "Bounded subprocess timeouts; rebuild retrieval/graph artifacts after successful acquisition; atomic JSON/JSONL writes; correct scheduler initialization and retry scheduling; standalone periodic worker/status artifact.", "No scheduled live refresh executed in this session; worker deployment remains open."),
        ("Exports", "backend/app/reporting/exporter.py; main.py", "Export exact saved answer; document-only dependency group; complete wrapped PDF text; distinguish query entities from real graph relationships.", "Prior Markdown/JSON tests passed; new DOCX/PDF content/regression checks remain."),
        ("Next.js client", "frontend/lib/api-client.ts; components/query-page.tsx; paper-workspace.tsx; app/papers/page.tsx", "Deduplicated refresh promises; streaming reset/error/incomplete-stream handling; request cancellation; saved-answer report downloads; paper upload/extraction/comparison; API helpers for jobs/history/saved evidence.", "TypeScript typecheck passed after dependency repair; production build status appears in Section 13."),
        ("Flutter client", "mobile/lib/services/api_service.dart; mobile/lib/main.dart", "Bounded SSE refresh retry; failure/reset handling; incomplete stream error; upload job polling and actual status; methods for notes/history/paper comparison/report export.", "Source edited; no Flutter/Dart SDK available, so compilation and device tests not claimed."),
        ("Environment/container fixes", "backend/requirements-documents.txt; backend/Dockerfile; frontend/Dockerfile; docker-compose.yml; .dockerignore; .gitignore; frontend/package.json", "Reused Python venv; ensured document dependencies; repaired Node installation on Linux filesystem; direct Node CLI scripts; Docker image includes scripts/ml/OCR and document packages; worker profile and build-context exclusions.", "Docker daemon unavailable. Windows mount chmod limitations documented; local Node symlink points to /tmp/omnirush/web-deps/node_modules."),
        ("Regression repair", "backend/tests/test_ingestion_security.py; test_new_features.py", "Corrected obsolete test that expected invalid image placeholder content; assert invalid bytes never become indexed evidence; optional FastAPI import handling; restore test runtime/settings and close worker resources.", "Full suite reached 36 passed / 1 skipped before later edits; latest captured verification below supersedes that checkpoint."),
    ]
    for title, files, work, status in changes:
        heading(title, 2); para("Files: " + files); para(work); para("Verification boundary: " + status)

    heading("6. Backend and database implementation detail")
    bullets(["The API supports a FastAPI/Uvicorn runtime plus a clearly labeled standard-library fallback. Both call the same route/domain services.", "Runtime rehydrates owned persisted documents, constructs authentication and retrieval services, connects optional infrastructure clients, and exposes detailed service health.", "Schema categories cover users/tokens, workspaces/research items/notes/saved evidence, documents/chunks/embeddings, conversations/messages/claims/citations, papers/gaps, cyber entities/relationships, experiments/metrics, ingestion jobs, audit/admin records, threats/vulnerabilities/IOCs, and newly stored answer snapshots.", "A declared schema table is not proof its complete workflow is used. Conversation/message/claim/citation tables and pgvector storage still need fuller service integration; current complete answers are preserved in the answers table.", "PostgreSQL/pgvector includes migration statements and a configured adapter. Redis and Neo4j include availability/health checks and lifecycle management. Live infrastructure verification is still outstanding.", "Local retrieval operates primarily over the in-memory document/chunk store and cached ranking structures. A persistent SQLite vector artifact exists from data bootstrap, but full runtime integration of that artifact/pgvector and stronger embeddings is not yet complete."])

    heading("7. Data acquisition, provenance and actual stored artifacts")
    bullets(["Official-source acquisition scripts exist for MITRE ATT&CK, MITRE ATLAS, NVD, CISA KEV, CWE, CAPEC, NIST and OWASP; OpenAlex acquisition provides permitted paper metadata and optional advertised full text.", "Fetches use bounded endpoints, caching, retry/backoff and source attribution. Failed fetches are recorded as failures rather than replaced by invented records.", "Normalization constructs stable source IDs, source aliases, normalized entities and relationships; unresolved relationships remain explicit.", "Processing builds sentence-aware chunks, BM25, deterministic CPU feature-hash embeddings, a persistent vector artifact and a graph artifact.", "The earlier final report records 8,931 normalized entities, 21,262 relationships, 9,455 chunks and 25 OpenAlex metadata records. These are historical recorded figures, not newly rerun counts."])
    manifest = read_json("data/data_manifest.json")
    source_rows = [(item.get("source", ""), item.get("status", ""), item.get("records", "not recorded"), item.get("version", ""), item.get("downloaded_at", item.get("processed_at", ""))) for item in manifest.get("sources", [])]
    para("Current manifest snapshot (fields read directly from data/data_manifest.json; a stored status is not a fresh network verification):")
    table(["Source", "Stored status", "Records", "Version", "Recorded time"], source_rows)
    table(["Stored artifact", "Recorded count", "Type/model/status"], [(item.get("name", ""), item.get("records", ""), item.get("index_type", item.get("embedding_model", item.get("status", "")))) for item in manifest.get("artifacts", [])])
    para(f"Manifest has {len(manifest.get('errors', []))} stored historical error entries. Historical failed attempts are not silently erased by this record.")

    heading("8. ML pipeline and research validity")
    bullets(["Intent and OOD baseline training scripts use TF-IDF plus Logistic Regression and write reproducible artifacts/metrics.", "Optional transformer training, inference and reranking entry points exist; installed packages alone do not mean a trained/loadable model exists.", "The earlier report records intent accuracy 0.0 on a tiny 14-item held-out set. This exposes the seed supervision weakness; it is not a useful production performance claim.", "No model accuracy improvement was measured in this resumed session. Training quality requires expanded attributed/human-labeled datasets, leakage checks, appropriate splits and evaluation.", "The existing development CyberTrust-QA file contains three items. It is suitable for contract smoke tests, not a broad scientific benchmark.", "The benchmark includes an intentionally unknown identifier to test abstention. That test question is not a factual vulnerability record.", "No new human annotations, papers, CVEs, research findings or benchmark scores have been generated to fill missing results."])
    heading("9. CERS: computation, interpretation and calibration")
    table(["Component", "Default weight", "Interpretation"], [("Source authority", "0.22", "Source metadata priority"), ("Relevance", "0.18", "Retrieval/query match"), ("Entailment", "0.25", "Claim support signal from active verifier"), ("Agreement", "0.15", "Conflict-adjusted cross-source signal"), ("Recency", "0.10", "Document publication/update freshness"), ("Retrieval confidence", "0.10", "Top normalized retrieval score")])
    para("Default raw-score abstention threshold: 0.42. Component weights can be configured and are normalized; new validation rejects invalid/non-finite/negative settings. CERS remains a project-defined indicator, not a validated probability of truth.")
    para("The revised calibration pipeline requires claim, evidence, source, category, human_label, annotator_id, finite cers_score and an explicit calibration/evaluation split. It rejects duplicate claim/evidence records, requires both supported and non-supported training cases, and defaults to at least 30 records per split. Human-label target: SUPPORTED versus not SUPPORTED. An isotonic mapping and threshold are fitted on the calibration split only; the held-out split supplies Brier score, reliability bins/ECE, acceptance coverage and selective risk.")
    para("No human calibration dataset has been supplied. Therefore no calibrated quality result is claimed. Runtime only attempts an optional artifact when CERS_CALIBRATION_PATH is configured, and the artifact must be fitted for the current weights and contain the label-file checksum.")
    heading("10. Evaluation and ablations")
    table(["Variant", "Current runner definition"], [("A", "Configured LLM without retrieval; retain generated outputs for scoring; unavailable if provider is not configured."), ("B", "Lexical retrieval baseline; other optional ranking/verification components disabled."), ("C", "Hybrid retrieval baseline."), ("D", "C plus cross-encoder reranking; NOT_RUN if a loadable reranker is unavailable."), ("E", "C plus source priority."), ("F", "C plus temporal ranking/timeline."), ("G", "C plus knowledge graph."), ("H", "C plus citation verification."), ("I", "C plus contradiction detection."), ("J", "All configured components, explicit fallbacks and CERS.")])
    para("D–I are independent additions to C, not a cumulative ladder. B–J share deterministic evidence generation so a changing external generator does not obscure component effects. Each result includes options, raw answer snapshots, retrieval metrics and runtime information. No absent component is reported as measured.")
    for filename in ("cybertrust_results.json", "ablation_results.json", "cers_calibration.json"):
        artifact = read_json("data/evaluation/" + filename)
        if not artifact: continue
        heading("Stored artifact: " + filename, 2)
        if filename == "ablation_results.json":
            table(["Variant", "Stored status", "Stored reason"], [(row.get("configuration", ""), row.get("status", ""), str(row.get("reason", ""))[:500]) for row in artifact.get("rows", [])])
            para("These statuses describe the stored artifact; the revised ablation implementation requires a new run before results are attributed to it.")
        else:
            summary = {key: value for key, value in artifact.items() if key not in {"items", "results", "rows"}}
            code(json.dumps(summary, indent=2, ensure_ascii=False)[:14000])

    heading("11. RAG security and ownership")
    bullets(["Password hashing, expiring signed tokens, persisted refresh-token rotation and authenticated APIs are implemented.", "Owner tags, repository lookups and route checks scope uploaded documents and research records to the authenticated owner.", "Prompt-injection patterns are recorded and flagged uploads are quarantined out of retrieval. This is a heuristic layer with potential false positives/negatives.", "Provider prompts delimit supplied evidence as untrusted data and require source citations; documents do not become system instructions.", "Citation marker validation and claim verification expose missing/unsupported evidence instead of merely attaching a citation-looking token.", "Log/IOC routines report local heuristic leads and explicitly avoid claiming external enrichment or a confirmed incident.", "Remaining security verification includes targeted authorization regressions on newly added routes, upload/resource boundary tests, broader RAG adversarial evaluation and production authentication/transport/deployment configuration."])

    heading("12. Environment, dependencies and deployment")
    table(["Component", "Observed situation / setup"], [("Operating environment", "Linux/WSL workspace at /mnt/c/Users/Lenovo/MLPA; Windows-mounted filesystem has chmod restrictions."), ("Python", "System Python 3.14.4 has no pip; existing /tmp/cybersecra-venv supplies working Python dependencies."), ("Document support", "Pillow, pytesseract, python-docx and reportlab present in the venv; PyMuPDF present. System Tesseract absent at the inspected checkpoint."), ("Node", "Node 24.21.0 and npm 11.19.0 available. Dependencies installed under /tmp/omnirush/web-deps after mount permission failures; frontend/node_modules points there locally."), ("Next.js", "Pinned Next 15.5.9; React 18.3.1; TypeScript 5.8.3. Typecheck succeeded after repair; build result is recorded below."), ("Docker", "Docker command/daemon unavailable in current environment; Compose configuration can still be statically checked."), ("Flutter / Dart", "SDKs unavailable at inspected checkpoint; mobile source is not build-verified."), ("Privilege", "sudo -n true requested interactive authentication, so unattended system package installation was unavailable."), ("External providers", "No configured live LLM/model results or human calibration results are asserted."), ("Git history", "main has no commits; working source is untracked. Per-command safe.directory was used without changing global Git configuration.")])
    heading("Repeatable commands", 2)
    for command in [
        "PYTHONPATH=backend /tmp/cybersecra-venv/bin/python backend/run.py",
        "PYTHONPATH=backend /tmp/cybersecra-venv/bin/python -m pytest -p no:cacheprovider -q backend/tests",
        "npm run dev --prefix frontend\nnpm run typecheck --prefix frontend\nnpm run build --prefix frontend",
        "PYTHONPATH=backend python scripts/bootstrap_data.py --minimal\nPYTHONPATH=backend python scripts/bootstrap_data.py --full",
        "PYTHONPATH=backend python scripts/evaluation/run_ablation.py\nPYTHONPATH=backend python scripts/evaluation/calibrate_cers.py --labels <human-reviewed-labels.jsonl>",
        "PYTHONPATH=backend python scripts/setup_models.py --nli\nPYTHONPATH=backend python scripts/sync_graph.py",
        "docker compose --profile infra --profile workers up --build",
        "flutter pub get\nflutter analyze\nflutter test\nflutter run --dart-define=API_BASE_URL=http://10.0.2.2:8000",
    ]: code(command)
    para("Commands referencing Docker, Flutter, model downloads or labels require those prerequisites. The temporary Python/Node paths describe this environment; a fresh machine should install equivalent dependencies in its own environment. Docker infra also needs matching DATABASE_URL/Neo4j credentials; enabling a Compose profile alone does not switch SQLite to PostgreSQL.")

    heading("13. Verification record")
    para("Earlier project report: 27 backend tests and 5 web tests passed; live Uvicorn health/auth/query checks and bootstrap runs were recorded. Those are prior-session claims from docs/FINAL_PROJECT_REPORT.md, not reruns by this document.")
    para("Current-session initial system-Python run found two errors: missing FastAPI import in one test module and an obsolete invalid-image expectation. The first venv run produced 35 passed, one failed, one skipped. After repairs, the suite reached 36 passed, one skipped; a later focused suite produced nine passed. Those checkpoints preceded some subsequent edits.")
    if checks:
        table(["Latest captured check", "Status", "Exit", "Finished (UTC)"], [(row["check"], row["status"], row["exit_code"], row["finished_at"]) for row in checks])
        para("A checker's exit status means the checker ran according to its own rules. Its output must be read before claiming service availability; e.g. a mobile checker can report SDKs absent. Full command output is included in Appendix E and saved alongside this document.")
    else:
        para("No embedded verification log was generated on this invocation. Use --verify to capture a current run.")

    heading("14. What is completed, partially complete, and still outstanding")
    table(["Area", "Current implementation", "Remaining work / completion evidence"], [
        ("Core backend", "Substantial local-first API/query/auth/research implementation; deterministic contract covered by existing tests.", "Targeted regression tests for the new snapshot, recovery, authorization, graph-depth, calibration and stream-failure paths; concurrency/load verification."),
        ("Database/retrieval infrastructure", "SQLite persistence, normalized schemas, optional PostgreSQL/pgvector/Redis/Neo4j adapters; answer/chunk APIs and graph importer.", "Start real services, run live repository parity tests, connect persistent vector indices/pgvector into query retrieval, measure scale and cache behavior."),
        ("Data pipeline", "Real acquisition, manifests, normalization, chunks/indices/graph and refresh worker source.", "Run revised refresh/reindex flow; verify atomic snapshot reload and failure recovery; deploy/supervise the scheduled worker."),
        ("ML models", "CPU baseline training/inference and optional transformer/reranker adapters.", "Install compatible ML stack/model weights, expand attributed labels, train/evaluate appropriately and report actual metrics."),
        ("Knowledge graph", "Local provenance graph, bounded traversal, optional real Neo4j import/query, retrieval links.", "Execute live Neo4j sync/search and bounded traversal tests; assess relation coverage and graph-assisted retrieval quality."),
        ("Research workspace", "Owned items, questions, notes, answers and saved evidence APIs.", "Complete frontend saved-evidence CRUD/history presentation and ensure all new cross-record references have tests."),
        ("Paper intelligence", "Section extraction, provenance, analysis/compare endpoints and paper workspace UI.", "Test varied real PDF/DOCX layouts, distinguish metadata-only records from full papers, improve extraction accuracy with evidence-backed parsing."),
        ("Multimodal ingestion", "Validated text/JSON/CSV/PDF/DOCX/image OCR adapters and job spool.", "Install system OCR, test real scanned/image documents and tables, stress limits/restart recovery. General visual understanding is not implemented."),
        ("Streaming", "Actual provider stream adapter, deterministic buffered fallback, completion verification, client resets/errors/cancel path.", "Exercise real configured provider, disconnections, slow model stages, network failures and authentication expiry; verify transport/cancellation under load."),
        ("Exports", "JSON/Markdown/DOCX/PDF code and exact saved-answer export.", "Verify new PDF wrapping, DOCX/PDF content preservation and multilingual rendering with real report fixtures."),
        ("Frontend", "Next.js routes, auth, query/evidence, papers, workspaces, resources and report download source.", "Resolve any captured build failures; perform browser end-to-end/accessibility checks and finish job/history/evidence UX."),
        ("Mobile", "Flutter auth/query/voice/camera/file/tools/research source; stricter SSE and status polling.", "Install SDK, generate/inspect platform projects, run analyze/tests/builds and device permissions/network/export/workspace flows."),
        ("RAG security", "Quarantine, ownership, citation checks and explicit heuristic scan/evaluation.", "Expand attributed adversarial evaluation; test false positives/negatives and every new ownership boundary."),
        ("Evaluation/ablation", "Revised independently switched runner and raw-result preservation.", "Run it, review metrics against the exact current code/corpus, enlarge human-reviewed benchmark and add uncertainty analysis."),
        ("CERS calibration", "Human-only validation/fitting/evaluation implementation.", "Collect/adjudicate human annotations; split calibration/evaluation; test fitter; fit/review artifact and report held-out results. No labels or metrics may be invented."),
        ("Deployment", "Dockerfiles, persistent services, worker profile and setup/check scripts.", "Docker live build/start/check, correct environment URLs/secrets, production worker/service supervision and backup/restore tests."),
    ])
    heading("Prioritized next actions", 2)
    bullets(["1. Complete targeted regression coverage for the new code, then rerun backend/API and web checks.", "2. Finish frontend production build and browser query/upload/paper/workspace/export tests.", "3. Run revised ablations and calibration validation with available real inputs; preserve NOT_RUN/PENDING states where prerequisites are absent.", "4. Verify real document formats and durable ingestion recovery; finish database/vector/graph integration tests against live services.", "5. Install/configure Docker, Flutter, OCR and optional model/provider prerequisites where feasible; deploy and test the refresh worker.", "6. Collect human-reviewed benchmark/calibration data and conduct the research evaluation before claiming calibrated reliability or benchmark superiority.", "7. Update older status documents and the saved plan so they reflect the final verified handoff; create a Git baseline only when explicitly requested."])

    heading("15. How to demonstrate and explain the project")
    bullets(["Start the API and Next.js client. Show /health/detailed so the audience sees the actual database, optional service availability and fallback modes.", "Register/log in. Ask ‘What does NIST CSF 2.0 provide?’ and inspect the final answer, evidence source links, claim status, limitations and CERS components.", "Explain that the stream shows real pipeline activity and, with a configured provider, actual generation; final verification determines the completed answer.", "Upload an attributed text/PDF paper, inspect processing status and source/chunk metadata, then compare selected papers and review missing fields/provenance.", "Use a workspace to save a question/answer and notes; inspect ownership with a second account during testing.", "Search a known MITRE technique in the graph and explain that nodes/edges come from the normalized source snapshot.", "Export a saved answer; verify its answer_id and citations match the viewed response.", "Open evaluation/environment views and discuss the small benchmark, optional models, unconfigured services and human-calibration requirement honestly."])
    heading("Five-minute explanation outline", 2)
    para("Minute 1: State the trust problem in cybersecurity research and show the overall architecture. Minute 2: Walk through one question from intent to ranked evidence. Minute 3: Explain claim verification, source disagreement, CERS and abstention. Minute 4: Show paper/workspace/graph/export workflows and their provenance. Minute 5: Present actual checks and clearly identify unfinished external integration, mobile and human-evaluation work.")
    heading("16. Glossary and important distinctions")
    table(["Term", "Meaning in this project"], [("RAG", "Retrieval-augmented generation: use retrieved source text to ground a response."), ("BM25", "Lexical retrieval scoring based on term frequency, document frequency and length."), ("TF-IDF / feature hash", "Lightweight numerical text representations used as explicit CPU baselines, distinct from pretrained semantic encoders."), ("RRF", "Reciprocal-rank fusion, combining ranked retrieval lists."), ("NLI", "Natural-language inference: compare premise/evidence with a hypothesis/claim."), ("OOD", "Out-of-domain query detection; avoids confidently answering unrelated questions."), ("CERS", "Project-defined Cyber Evidence Reliability Score; not automatically a probability of truth."), ("Provenance", "Where evidence came from: source, URL, record/chunk/page/section and timestamps/version where available."), ("Abstention", "Explicitly declining to give a confident answer when evidence is inadequate."), ("Ablation", "Run the system with controlled component changes to measure their effect."), ("Calibration", "Fit/assess a score against independent human labels; requires held-out evaluation."), ("Implemented / verified", "Source functionality exists / the specified checks or live workflow actually executed successfully."), ("Quarantined", "Retained document with security flags that is excluded from retrieval."), ("SSE", "Server-sent events over HTTP for pipeline status, answer deltas/resets, errors and completion.")])

    doc.add_page_break()
    heading("Appendix A. API route inventory")
    main_tree = ast.parse((ROOT / "backend/app/main.py").read_text(encoding="utf-8"))
    routes = []
    for node in ast.walk(main_tree):
        if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)): continue
        for decorator in node.decorator_list:
            if isinstance(decorator, ast.Call) and isinstance(decorator.func, ast.Attribute) and decorator.args and isinstance(decorator.args[0], ast.Constant) and isinstance(decorator.args[0].value, str) and decorator.args[0].value.startswith("/"):
                method = decorator.func.attr.upper()
                for keyword in decorator.keywords:
                    if keyword.arg == "methods": method = ", ".join(ast.literal_eval(keyword.value))
                routes.append((method, decorator.args[0].value, node.name, f"main.py:{node.lineno}"))
    table(["Method", "Path", "Handler", "Source"], sorted(routes, key=lambda row: row[1]))
    para("Inventory is extracted from FastAPI route declarations. The common _route dispatcher and standard-library fallback use related paths; exact request bodies and authorization are implemented in backend/app/main.py.")

    heading("Appendix B. Database table inventory")
    sqlite_text = (ROOT / "backend/app/database/sqlite.py").read_text(encoding="utf-8")
    tables = sorted(set(re.findall(r"CREATE TABLE IF NOT EXISTS\s+(\w+)", sqlite_text)))
    table(["SQLite tables declared in schema"], [(name,) for name in tables])
    para("PostgreSQL migrations additionally declare vector_embeddings and pgvector/HNSW structures. Tables are listed as schema implementation, not as proof every product workflow uses every table.")

    heading("Appendix C. Granular Python class/function inventory")
    para("Every named class/function/method below is extracted from the current Python source, with line numbers and source docstrings when present. Test helpers are included in the test modules. A symbol inventory does not prove that every function was executed or externally verified. Names without a docstring are deliberately not assigned an invented behavior description.")
    files = inventory()
    inventory_manifest = []
    symbol_count = 0
    for path in files:
        relative = str(path.relative_to(ROOT))
        inventory_manifest.append({"path": relative, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
        if path.suffix != ".py": continue
        try: module_doc, definitions = python_definitions(path)
        except SyntaxError as exc:
            heading(relative, 2); para(f"Source parse failed: {exc}"); continue
        if not definitions and not module_doc: continue
        heading(relative, 2)
        if module_doc: para(module_doc.split("\n")[0])
        for line, kind, signature, description in definitions:
            symbol_count += 1
            code(f"L{line} [{kind}] {signature}")
            if description: para(description)
        if not definitions: para("Module provides imports/constants/data; no named Python functions/classes.")

    heading("Appendix D. Frontend/mobile modules and entry points")
    for path in files:
        if path.suffix == ".py": continue
        relative = str(path.relative_to(ROOT))
        text = path.read_text(encoding="utf-8")
        names = re.findall(r"(?:export\s+)?(?:async\s+)?function\s+(\w+)|(?:export\s+)?class\s+(\w+)|class\s+(\w+)\s+extends", text)
        names = sorted({next(part for part in row if part) for row in names})
        heading(relative, 2)
        para("Named entry points: " + (", ".join(names) if names else "Route/component/type/test/helper module; inspect the source for inline handlers and declarations."))
        if path.suffix == ".dart" and "api_service" in relative:
            methods = re.findall(r"(?:Future<[^\n]+?>|static\s+\w+|void)\s+(_?\w+)\s*\(", text)
            para("Detected service methods: " + ", ".join(sorted(set(methods))))
    para(f"Source snapshot covers {len(files)} code files and {symbol_count} named Python classes/functions/methods. File hashes are saved in work_record_source_manifest.json to identify the documented snapshot.")

    heading("Appendix E. Exact verification command output")
    for row in checks:
        heading(row["check"], 2)
        code(row["command"])
        para(f"Status: {row['status']} | Exit: {row['exit_code']} | UTC: {row['started_at']} → {row['finished_at']}")
        # Keep logs complete; long environment JSON is intentionally retained.
        code(row["output"] or "No output; consult exit status.")

    heading("Appendix F. Sources for this work record")
    bullets(["__agent__/plan.json — saved task checkpoint and timestamp.", "CHANGELOG.md — initial increment and scope notes.", "README.md — overall architecture, commands and feature descriptions; some original environment wording is stale.", "docs/FINAL_PROJECT_REPORT.md — earlier implementation, bootstrap counts and verification claims.", "docs/QUALITY_GATE.md — earlier completion/verification checklist.", "docs/DECISIONS.md — local-first/fallback/schema/trust-boundary rationale.", "docs/datasets/bootstrapping.md; docs/research/{methodology,limitations,ablation}.md — data/model/evaluation methodology and original boundaries.", "data/data_manifest.json; data/data_version.json; data/evaluation/*.json — stored run artifacts, read as stored evidence rather than fresh measurements.", "Current source files listed in the manifest — actual implementation snapshot.", "docs/reports/work_record_verification.json — exact checks captured during document generation."])
    para("This is a living work record. Further implementation will change line numbers and statuses. Regenerate it after subsequent verified changes; do not reinterpret this snapshot as a timestamped transcript of actions that were never recorded.")

    path = OUTPUT / "CyberSecRA-X_Complete_Work_Record_2026-10-03.docx"
    doc.save(path)
    (OUTPUT / "work_record_source_manifest.json").write_text(json.dumps({"generated_at": generated_at, "files": inventory_manifest, "python_symbols": symbol_count}, indent=2), encoding="utf-8")
    with ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise RuntimeError("DOCX ZIP integrity check failed")
    restored = Document(path)
    if len(restored.paragraphs) < 100 or len(restored.tables) < 10:
        raise RuntimeError("Generated work record is unexpectedly incomplete")
    print(json.dumps({"document": str(path), "bytes": path.stat().st_size, "paragraphs": len(restored.paragraphs), "tables": len(restored.tables), "source_files": len(files), "python_symbols": symbol_count, "docx_integrity": "passed"}, indent=2))


if __name__ == "__main__":
    main()
